"""
ppt_writer.py
=============
Generate an on-brand PowerPoint (.pptx) summary deck from the comparison
`sheets` produced by compare_engine.compare_renewals().

The deck uses the bundled Cisco template (assets/cisco_template.pptx) so every
slide inherits the Cisco theme, fonts and colours. Seven slides are produced:

  1. Cover
  2. Executive Summary (3 KPI cards)
  3. Value Bridge - Service (waterfall + detail table)
  4. Value Bridge - SW Subscription (waterfall + detail table)
  5. Added Items (table)
  6. Removed Items (table: LDOS vs customer)
  7. SW Subscription Changes (counts + product detail)

Charts are drawn as native shapes (rectangles / connectors) so the result is
100% reliable on both Mac and Windows with no external chart engine.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ── Cisco brand palette ────────────────────────────────────────────────────────
DARK  = RGBColor(0x0D, 0x27, 0x4D)   # Cisco midnight blue
BLUE  = RGBColor(0x00, 0xBC, 0xEB)   # Cisco blue
GREEN = RGBColor(0x6C, 0xC0, 0x4A)   # positive / add
AMBER = RGBColor(0xFB, 0xAB, 0x18)   # removal
RED   = RGBColor(0xE2, 0x23, 0x1A)   # negative net
GREY  = RGBColor(0x58, 0x59, 0x5B)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

TEMPLATE_PATH = Path(__file__).parent / "assets" / "cisco_template.pptx"

# Layout indices in the Cisco template
LAYOUT_TITLE_ONLY = 9    # 'Title Only 1'  (title + footer + slide number)
LAYOUT_COVER      = 10   # 'Title, Subtitle Only 1'


# ── small helpers ──────────────────────────────────────────────────────────────
def _money(v: float) -> str:
    v = float(v or 0)
    if v < 0:
        return f"-${abs(v):,.0f}"
    return f"${v:,.0f}"


def _num(row, col: str) -> float:
    if row is None:
        return 0.0
    v = row.get(col, "")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return 0.0


def _txt(row, col: str) -> str:
    if row is None:
        return ""
    v = row.get(col, "")
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("", "nan", "none") else s


def _parse_summary(summary: pd.DataFrame):
    """Return (top, child): top[letter]=row for A..G; child[(parent,sec)]=row."""
    top: dict[str, pd.Series] = {}
    child: dict[tuple[str, str], pd.Series] = {}
    parent = None
    for _, r in summary.iterrows():
        sec = str(r.get("section", "") or "").strip()
        if sec in list("ABCDEFG"):
            parent = sec
            top[sec] = r
        elif sec and sec != "NOTE" and parent is not None:
            child[(parent, sec)] = r
    return top, child


def _delete_all_slides(prs: Presentation) -> None:
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass
        sldIdLst.remove(sldId)


def _add_slide(prs: Presentation, layout_idx: int):
    layout = prs.slide_layouts[layout_idx]
    return prs.slides.add_slide(layout)


def _set_title(slide, text: str) -> None:
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text = text
            for p in ph.text_frame.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = DARK
            return
    # fallback: draw our own title
    _textbox(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.9),
             text, size=28, bold=True, color=DARK)


def _textbox(slide, x, y, w, h, text, size=14, bold=False, color=DARK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def _rect(slide, x, y, w, h, fill, line=None, line_w=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(1)
    shp.shadow.inherit = False
    return shp


# ── waterfall chart drawn as shapes ────────────────────────────────────────────
def _draw_waterfall(slide, x, y, w, h, bars, currency="$"):
    """bars: list of dicts {label, lo, hi, kind, end}.
    kind in {'total','up','down'}; end = running cumulative used for connectors.
    """
    x, y, w, h = int(x), int(y), int(w), int(h)
    lows = [min(b["lo"], b["hi"]) for b in bars]
    highs = [max(b["lo"], b["hi"]) for b in bars]
    vmax = max(highs + [0])
    vmin = min(lows + [0])
    span = (vmax - vmin) or 1.0

    plot_top = y
    plot_h = int(h * 0.78)          # leave room for category labels
    label_band = h - plot_h

    def yv(val):  # value -> EMU y (top-down)
        frac = (vmax - val) / span
        return plot_top + int(frac * plot_h)

    n = len(bars)
    gap = int(w * 0.04)
    bar_w = int((w - gap * (n + 1)) / n)

    colors = {"total": DARK, "up": GREEN, "down": AMBER}

    prev_cx_right = None
    prev_end_y = None
    for i, b in enumerate(bars):
        bx = x + gap + i * (bar_w + gap)
        lo, hi = min(b["lo"], b["hi"]), max(b["lo"], b["hi"])
        top_y = yv(hi)
        bot_y = yv(lo)
        bh = max(bot_y - top_y, Emu(Pt(2)))
        _rect(slide, bx, top_y, bar_w, bh, colors.get(b["kind"], DARK))

        # value label above the bar
        if b["kind"] == "total":
            lbl = _money(hi)
        else:
            delta = b["hi"] - b["lo"] if b["kind"] == "up" else -(hi - lo)
            lbl = ("+" if delta >= 0 else "-") + _money(abs(delta)).lstrip("-")
        _textbox(slide, bx - gap // 2, top_y - Inches(0.32), bar_w + gap, Inches(0.3),
                 lbl, size=10, bold=True, color=DARK, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.BOTTOM)

        # category label below plot
        _textbox(slide, bx - gap // 2, plot_top + plot_h + Emu(Pt(2)),
                 bar_w + gap, label_band, b["label"], size=9, bold=False,
                 color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)

        # connector from previous bar's running-total level
        end_y = yv(b["end"])
        if prev_cx_right is not None and prev_end_y is not None:
            conn = slide.shapes.add_connector(2, prev_cx_right, prev_end_y, bx, prev_end_y)
            conn.line.color.rgb = GREY
            conn.line.width = Pt(0.75)
            conn.line.dash_style = None
        prev_cx_right = bx + bar_w
        prev_end_y = end_y


# ── styled table ───────────────────────────────────────────────────────────────
def _add_table(slide, x, y, w, headers, rows, col_widths=None,
               font_size=10, header_fill=DARK, max_rows=18):
    nrow = min(len(rows), max_rows) + 1
    ncol = len(headers)
    truncated = len(rows) > max_rows
    height = Inches(0.34) * nrow
    gfx = slide.shapes.add_table(nrow, ncol, int(x), int(y), int(w), int(height))
    table = gfx.table

    if col_widths:
        total = sum(col_widths)
        for c, cw in enumerate(col_widths):
            table.columns[c].width = int(w * cw / total)

    # header
    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(htext)
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = WHITE

    # body
    for r_i in range(min(len(rows), max_rows)):
        row = rows[r_i]
        fill = WHITE if r_i % 2 == 0 else LIGHT
        for c in range(ncol):
            cell = table.cell(r_i + 1, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.01)
            cell.margin_bottom = Inches(0.01)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(row[c])
            run.font.size = Pt(font_size)
            run.font.color.rgb = DARK
    if truncated:
        _textbox(slide, x, y + height + Inches(0.02), w, Inches(0.3),
                 f"…and {len(rows) - max_rows} more (see Excel report)",
                 size=9, italic=True, color=GREY)
    return gfx


# ── individual slides ──────────────────────────────────────────────────────────
def _slide_cover(prs, customer, period, as_of):
    slide = _add_slide(prs, LAYOUT_COVER)
    _set_title(slide, "Renewal Quote Comparison")
    sub = customer or "Customer"
    if period:
        sub += f"\n{period}"
    # subtitle placeholder (idx 11) if present
    filled = False
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 11:
            ph.text = sub
            filled = True
            break
    if not filled:
        _textbox(slide, Inches(0.7), Inches(3.4), Inches(11), Inches(1.5),
                 sub, size=20, bold=False, color=GREY)
    _textbox(slide, Inches(0.7), Inches(6.6), Inches(11), Inches(0.5),
             f"Generated {as_of:%d %b %Y}", size=12, color=GREY)
    return slide


def _slide_exec_summary(prs, top):
    slide = _add_slide(prs, LAYOUT_TITLE_ONLY)
    _set_title(slide, "Executive Summary")
    last = _num(top.get("A"), "Total $")
    this = _num(top.get("B"), "Total $")
    net = _num(top.get("C"), "Total $")
    pct = (net / last * 100) if last else 0.0

    cards = [
        ("Last Year", _money(last), DARK),
        ("This Year", _money(this), BLUE),
        ("Net Change", ("+" if net >= 0 else "-") + _money(abs(net)).lstrip("-")
         + f"  ({pct:+.1f}%)", GREEN if net >= 0 else RED),
    ]
    cw = Inches(3.9)
    gap = Inches(0.35)
    total_w = cw * 3 + gap * 2
    start_x = (prs.slide_width - total_w) // 2
    top_y = Inches(2.3)
    ch = Inches(2.2)
    for i, (label, value, color) in enumerate(cards):
        cx = start_x + i * (cw + gap)
        _rect(slide, cx, top_y, cw, ch, LIGHT)
        _rect(slide, cx, top_y, cw, Inches(0.14), color)   # accent stripe
        _textbox(slide, cx, top_y + Inches(0.5), cw, Inches(0.5), label,
                 size=16, bold=True, color=GREY, align=PP_ALIGN.CENTER)
        _textbox(slide, cx, top_y + Inches(1.05), cw, Inches(0.9), value,
                 size=24, bold=True, color=color, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)

    _textbox(slide, start_x, top_y + ch + Inches(0.5), total_w, Inches(1.2),
             "Total = calculated extended list price across Technical Services (TS) "
             "and Software Subscriptions (Sub-TnC). Net change reconciles exactly as "
             "This Year = Last Year + Additions - Removals.",
             size=12, color=GREY, align=PP_ALIGN.CENTER)
    return slide


def _waterfall_bars(A, D, E, B):
    return [
        {"label": "Last Year", "lo": 0, "hi": A, "kind": "total", "end": A},
        {"label": "Add Value", "lo": A, "hi": A + D, "kind": "up", "end": A + D},
        {"label": "Removal Value", "lo": A + D, "hi": A + D + E, "kind": "down", "end": A + D + E},
        {"label": "This Year", "lo": 0, "hi": B, "kind": "total", "end": B},
    ]


def _slide_bridge(prs, title, A, D, E, B, add_rows, remove_rows):
    slide = _add_slide(prs, LAYOUT_TITLE_ONLY)
    _set_title(slide, title)
    bars = _waterfall_bars(A, D, E, B)
    _draw_waterfall(slide, Inches(0.7), Inches(1.6), Inches(7.4), Inches(4.6), bars)

    # ── Two grouped tables on the right: Add (green) then Remove (orange) ──────
    tx = Inches(8.5)
    tw = Inches(4.2)
    row_h = Inches(0.30)
    cw = [2.6, 1.4]

    def _write_grouped_table(y_top, group_label, rows_data, total_val, hdr_color):
        """Draw a header row + data rows + total row; return bottom y."""
        nrow = len(rows_data) + 2   # header + data + total
        height = row_h * nrow
        gfx = slide.shapes.add_table(nrow, 2,
                                      int(tx), int(y_top), int(tw), int(height))
        t = gfx.table
        total_cw = sum(cw)
        for c, w in enumerate(cw):
            t.columns[c].width = int(tw * w / total_cw)

        def _fmt_cell(cell, text, bold, fill_rgb, font_color, align_right=False, font_size=9):
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill_rgb
            cell.margin_left = Inches(0.05); cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if align_right else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(text)
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = font_color

        # Header row
        _fmt_cell(t.cell(0, 0), group_label, True, hdr_color, WHITE)
        _fmt_cell(t.cell(0, 1), "", True, hdr_color, WHITE)

        # Data rows
        for i, (label, value) in enumerate(rows_data, start=1):
            fill = WHITE if i % 2 == 0 else LIGHT
            _fmt_cell(t.cell(i, 0), label, False, fill, DARK)
            _fmt_cell(t.cell(i, 1), value, False, fill, DARK, align_right=True)

        # Total row
        total_fill = RGBColor(0xE8, 0xF5, 0xE1) if hdr_color == GREEN else RGBColor(0xFF, 0xF0, 0xD0)
        total_font = GREEN if hdr_color == GREEN else AMBER
        last = nrow - 1
        _fmt_cell(t.cell(last, 0), "TOTAL", True, total_fill, total_font)
        _fmt_cell(t.cell(last, 1), total_val, True, total_fill, total_font, align_right=True)

        return y_top + height + Inches(0.10)

    add_total  = sum(float(str(v).replace("$","").replace(",","").replace("+","")) for _, v in add_rows if str(v).replace("$","").replace(",","").replace("+","").lstrip("-").replace(".","").isdigit() or "$" in str(v))
    rem_total  = sum(float(str(v).replace("$","").replace(",","").replace("+","")) for _, v in remove_rows if str(v).replace("$","").replace(",","").replace("+","").lstrip("-").replace(".","").isdigit() or "$" in str(v))

    def _parse_money(rows):
        total = 0.0
        for _, v in rows:
            s = str(v).replace("$","").replace(",","").replace("+","").strip()
            try: total += float(s)
            except ValueError: pass
        return _money(total)

    y = Inches(1.7)
    y = _write_grouped_table(y, "▲  Add Value", add_rows, _parse_money(add_rows), GREEN)
    _write_grouped_table(y, "▼  Removal Value", remove_rows, _parse_money(remove_rows), AMBER)
    return slide


def _slide_added(prs, child):
    slide = _add_slide(prs, LAYOUT_TITLE_ONLY)
    _set_title(slide, "Added Items")
    rows = []
    for (parent, sec), r in child.items():
        if parent == "G" and sec.startswith("1.") and sec.count(".") == 1:
            rows.append([
                _txt(r, "metric"),
                _txt(r, "Total $"),                    # item type
                f"{_num(r, 'value svc ($)'):,.0f}",    # QTY (lines)
                f"{_num(r, 'value SW ($)'):,.0f}",      # QTY sum
            ])
    if not rows:
        _textbox(slide, Inches(0.7), Inches(2.5), Inches(11), Inches(1),
                 "No added price-bearing items in this comparison.", size=14, color=GREY)
        return slide
    _add_table(slide, Inches(0.7), Inches(1.6), Inches(11.9),
               ["Product ID", "Item Type", "Lines", "Qty Sum"], rows,
               col_widths=[4.5, 2.5, 1.5, 1.5], font_size=10, max_rows=16)
    return slide


def _slide_removed(prs, child):
    slide = _add_slide(prs, LAYOUT_TITLE_ONLY)
    _set_title(slide, "Removed Items")
    ldos_rows, cust_rows = [], []
    for (parent, sec), r in child.items():
        if parent != "G":
            continue
        if sec.startswith("2.1.") and sec.count(".") == 2:
            ldos_rows.append([_txt(r, "metric"),
                              f"{_num(r, 'value svc ($)'):,.0f}",
                              _txt(r, "date")])
        elif sec.startswith("2.2.") and sec.count(".") == 2:
            cust_rows.append([_txt(r, "metric"),
                              f"{_num(r, 'value svc ($)'):,.0f}"])
    if not ldos_rows and not cust_rows:
        _textbox(slide, Inches(0.7), Inches(2.5), Inches(11), Inches(1),
                 "No removed price-bearing items in this comparison.", size=14, color=GREY)
        return slide
    ldos_rows.sort(key=lambda x: x[0])
    cust_rows.sort(key=lambda x: x[0])

    _textbox(slide, Inches(0.7), Inches(1.5), Inches(6.0), Inches(0.35),
             "Removed due to LDOS (end of support)", size=13, bold=True, color=AMBER)
    _add_table(slide, Inches(0.7), Inches(1.95), Inches(6.0),
               ["Product ID", "QTY", "LDOS Date"], ldos_rows,
               col_widths=[3.4, 1.0, 1.8], font_size=9, max_rows=14)

    _textbox(slide, Inches(7.1), Inches(1.5), Inches(5.5), Inches(0.35),
             "Removed by customer", size=13, bold=True, color=GREY)
    _add_table(slide, Inches(7.1), Inches(1.95), Inches(5.5),
               ["Product ID", "QTY"], cust_rows,
               col_widths=[4.0, 1.2], font_size=9, max_rows=14)
    return slide


def _slide_sw_changes(prs, top, child):
    slide = _add_slide(prs, LAYOUT_TITLE_ONLY)
    _set_title(slide, "SW Subscription Changes")

    # counts from F.6-F.9
    def fcount(sec):
        r = child.get(("F", sec))
        return int(_num(r, "value svc ($)")) if r is not None else 0

    summary_rows = [
        ["Added (Sub-TnC)", str(fcount("6"))],
        ["Potential expired & re-bought", str(fcount("7"))],
        ["Removed (Sub-TnC)", str(fcount("8"))],
        ["Tier changes", str(fcount("9"))],
    ]
    _add_table(slide, Inches(0.7), Inches(1.6), Inches(4.6),
               ["Change type", "Count"], summary_rows,
               col_widths=[3.2, 1.0], font_size=11, max_rows=8)

    # product detail from G.4 / G.5 / G.6 children
    detail = []
    labelmap = {"4": "Add", "5": "Remove", "6": "Tier change"}
    for (parent, sec), r in child.items():
        if parent != "G":
            continue
        head = sec.split(".")[0]
        if head in labelmap and sec.count(".") >= 1:
            product = _txt(r, "metric")
            if not product:
                continue
            qty = _num(r, "value svc ($)")
            detail.append([product, labelmap[head], f"{qty:,.0f}"])
    if detail:
        _add_table(slide, Inches(5.7), Inches(1.6), Inches(6.9),
                   ["Product ID", "Change", "QTY"], detail,
                   col_widths=[4.0, 1.6, 1.2], font_size=10, max_rows=16)
    else:
        _textbox(slide, Inches(5.7), Inches(1.9), Inches(6.9), Inches(1),
                 "No individual SW subscription line changes detected.",
                 size=12, color=GREY)
    return slide


# ── public entry point ─────────────────────────────────────────────────────────
def write_ppt(sheets: dict[str, pd.DataFrame], customer_name: str = "",
              period_label: str = "", as_of: date | None = None) -> bytes:
    as_of = as_of or date.today()
    summary = sheets.get("Summary")
    if summary is None or summary.empty:
        raise ValueError("Summary sheet is empty; run the comparison first.")

    top, child = _parse_summary(summary)

    prs = Presentation(str(TEMPLATE_PATH))
    _delete_all_slides(prs)

    # 1. Cover
    _slide_cover(prs, customer_name, period_label, as_of)
    # 2. Executive Summary
    _slide_exec_summary(prs, top)
    # 3. Value Bridge - Service
    A, D, E, B = (_num(top.get(k), "value svc ($)") for k in ("A", "D", "E", "B"))
    # Tier/SLA split is optional (only present when compare_renewals() was given
    # egrid_lookup/glasia_lookup) -- fall back to the combined row otherwise.
    # Split mode renumbers: 6.2=Tier, 6.3=SLA, 6.4=Unit Price, 6.5=QTY.
    tier_sla_split = ("D", "6.5") in child
    if tier_sla_split:
        sla_tier_add_rows = [
            ["Higher Tier", _money(_num(child.get(("D", "6.2")), "value svc ($)"))],
            ["Higher SLA",  _money(_num(child.get(("D", "6.3")), "value svc ($)"))],
        ]
        sla_tier_rem_rows = [
            ["Lower Tier",  _money(_num(child.get(("E", "6.2")), "value svc ($)"))],
            ["Lower SLA",   _money(_num(child.get(("E", "6.3")), "value svc ($)"))],
        ]
        unit_price_sec, qty_sec = "6.4", "6.5"
    else:
        sla_tier_add_rows = [["Higher SLA / Tier", _money(_num(child.get(("D", "6.2")), "value svc ($)"))]]
        sla_tier_rem_rows = [["Lower SLA / Tier",   _money(_num(child.get(("E", "6.2")), "value svc ($)"))]]
        unit_price_sec, qty_sec = "6.3", "6.4"
    svc_add = [
        ["TS - Add SN#",      _money(_num(child.get(("D", "1")),   "value svc ($)"))],
        ["TS - Add Instance#",_money(_num(child.get(("D", "2")),   "value svc ($)"))],
        ["Longer Duration",   _money(_num(child.get(("D", "6.1")), "value svc ($)"))],
        *sla_tier_add_rows,
        ["Higher Unit Price", _money(_num(child.get(("D", unit_price_sec)), "value svc ($)"))],
        ["Higher QTY",        _money(_num(child.get(("D", qty_sec)),        "value svc ($)"))],
    ]
    svc_rem = [
        ["TS - Remove SN#",      _money(_num(child.get(("E", "1")),   "value svc ($)"))],
        ["TS - Remove Instance#",_money(_num(child.get(("E", "2")),   "value svc ($)"))],
        ["Shorter Duration",     _money(_num(child.get(("E", "6.1")), "value svc ($)"))],
        *sla_tier_rem_rows,
        ["Lower Unit Price",     _money(_num(child.get(("E", unit_price_sec)), "value svc ($)"))],
        ["Lower QTY",            _money(_num(child.get(("E", qty_sec)),        "value svc ($)"))],
    ]
    _slide_bridge(prs, "Value Bridge - Service (TS)", A, D, E, B, svc_add, svc_rem)
    # 4. Value Bridge - SW Subscription
    A2, D2, E2, B2 = (_num(top.get(k), "value SW ($)") for k in ("A", "D", "E", "B"))
    sw_add = [
        ["Sub-TnC Add (key)",   _money(_num(child.get(("D", "3")),   "value SW ($)"))],
        ["Sub-TnC Add (new)",   _money(_num(child.get(("D", "4")),   "value SW ($)"))],
        ["Potential expired",   _money(_num(child.get(("D", "5")),   "value SW ($)"))],
        ["Longer Duration",     _money(_num(child.get(("D", "6.1")), "value SW ($)"))],
        ["Higher Tier",         _money(_num(child.get(("D", "6.2")), "value SW ($)"))],
    ]
    sw_rem = [
        ["Sub-TnC Remove (key)",  _money(_num(child.get(("E", "3")),   "value SW ($)"))],
        ["Sub-TnC Remove (new)",  _money(_num(child.get(("E", "4")),   "value SW ($)"))],
        ["Sub-TnC Remove (data)", _money(_num(child.get(("E", "5")),   "value SW ($)"))],
        ["Shorter Duration",      _money(_num(child.get(("E", "6.1")), "value SW ($)"))],
        ["Lower Tier",            _money(_num(child.get(("E", "6.2")), "value SW ($)"))],
    ]
    _slide_bridge(prs, "Value Bridge - SW Subscription (Sub-TnC)", A2, D2, E2, B2, sw_add, sw_rem)
    # 5. Added Items
    _slide_added(prs, child)
    # 6. Removed Items
    _slide_removed(prs, child)
    # 7. SW Subscription Changes
    _slide_sw_changes(prs, top, child)

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
